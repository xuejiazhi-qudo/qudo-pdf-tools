"""格式转换：PDF -> Word / Excel。"""

from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence

import pymupdf

from ..worker import Progress
from .pages import _open, _safe_dir

Progress_ = Optional[Progress]

WORD_EDITABLE = "editable"     # 真实可编辑文字 + 表格/版式（pdf2docx）
WORD_TEXT = "text"             # 纯文字段落（python-docx 自建）
WORD_FIDELITY = "fidelity"     # 整页图片，版式 100% 还原

WORD_MODE_LABELS = {
    WORD_EDITABLE: "可编辑（推荐）",
    WORD_TEXT: "纯文字",
    WORD_FIDELITY: "版式还原（图片）",
}

EXCEL_XLSX = "xlsx"
EXCEL_CSV = "csv"

SHEET_PER_TABLE = "per-table"
SHEET_PER_PAGE = "per-page"
SHEET_SINGLE = "single"


# ====================================================================== Word


def pdf_to_word(
    src_path: str,
    output: str,
    mode: str = WORD_EDITABLE,
    pages: Optional[Sequence[int]] = None,
    password: Optional[str] = None,
    progress: Progress_ = None,
) -> str:
    pr = progress or Progress()
    _safe_dir(output)

    if mode == WORD_EDITABLE:
        _word_editable(src_path, output, pages, password, pr)
    elif mode == WORD_TEXT:
        _word_text(src_path, output, pages, password, pr)
    elif mode == WORD_FIDELITY:
        _word_fidelity(src_path, output, pages, password, pr)
    else:
        raise ValueError(f"未知的 Word 转换模式：{mode}")

    pr.report(100, "完成")
    return output


def _word_editable(
    src_path: str,
    output: str,
    pages: Optional[Sequence[int]],
    password: Optional[str],
    pr: Progress,
) -> None:
    pr.report(10, "解析版面")
    try:
        from pdf2docx import Converter
    except Exception:
        # 打包版排除了 cv2/onnxruntime，pdf2docx 无法导入，走原生实现
        _word_editable_native(src_path, output, pages, password, pr)
        return

    conv = Converter(src_path, password=password or "")
    try:
        if pages is not None and len(pages):
            contiguous = _contiguous(pages)
            if contiguous:
                pr.report(30, "转换中")
                conv.convert(output, start=contiguous[0], end=contiguous[1] + 1)
            else:
                # 非连续页：先抽页到临时文档再转换
                import tempfile

                with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tf:
                    tmp = tf.name
                try:
                    from .pages import extract_pages

                    extract_pages(src_path, pages, tmp, password)
                    c2 = Converter(tmp)
                    try:
                        pr.report(30, "转换中")
                        c2.convert(output)
                    finally:
                        c2.close()
                finally:
                    if os.path.exists(tmp):
                        os.remove(tmp)
        else:
            pr.report(30, "转换中")
            conv.convert(output)
    finally:
        conv.close()


def _contiguous(pages: Sequence[int]) -> Optional[tuple]:
    ordered = sorted(set(pages))
    if not ordered:
        return None
    if ordered[-1] - ordered[0] + 1 == len(ordered):
        return (ordered[0], ordered[-1])
    return None


def _word_editable_native(
    src_path: str,
    output: str,
    pages: Optional[Sequence[int]],
    password: Optional[str],
    pr: Progress,
) -> None:
    """不依赖 pdf2docx 的可编辑导出：文本块 + 表格按阅读顺序写入 DOCX。"""
    import docx
    from docx.shared import Pt

    doc = _open(src_path, password)
    try:
        targets = list(pages) if pages else list(range(doc.page_count))
        out = docx.Document()
        out.styles["Normal"].font.size = Pt(11)

        total = len(targets)
        for k, i in enumerate(targets):
            pr.step(k, total, f"第 {i + 1} 页")
            if not 0 <= i < doc.page_count:
                continue
            if k:
                out.add_page_break()
            page = doc.load_page(i)

            tables = []
            try:
                for t in page.find_tables().tables or []:
                    rows = [[_clean_cell(c) for c in row] for row in (t.extract() or [])]
                    if len(rows) >= 2:
                        tables.append({"bbox": pymupdf.Rect(t.bbox), "rows": rows, "done": False})
            except Exception:  # noqa: BLE001
                tables = []

            blocks = [
                b for b in page.get_text("blocks")
                if b[6] == 0 and (b[4] or "").strip()
            ]

            def _inside(block_bbox) -> Optional[dict]:
                bb = pymupdf.Rect(block_bbox[:4])
                for tb in tables:
                    if tb["bbox"].contains(bb) or (bb & tb["bbox"]).get_area() > bb.get_area() * 0.6:
                        return tb
                return None

            items = sorted(blocks, key=lambda b: (round(b[1], 1), b[0]))
            for block in items:
                holder = _inside(block)
                if holder is not None:
                    if not holder["done"]:
                        holder["done"] = True
                        rows = holder["rows"]
                        cols = max(len(r) for r in rows)
                        table = out.add_table(rows=len(rows), cols=cols)
                        table.style = "Table Grid"
                        for r, row in enumerate(rows):
                            for c in range(cols):
                                table.cell(r, c).text = row[c] if c < len(row) else ""
                    continue
                text = " ".join(block[4].split())
                if text:
                    out.add_paragraph(text)

            for tb in tables:
                if not tb["done"]:
                    rows = tb["rows"]
                    cols = max(len(r) for r in rows)
                    table = out.add_table(rows=len(rows), cols=cols)
                    table.style = "Table Grid"
                    for r, row in enumerate(rows):
                        for c in range(cols):
                            table.cell(r, c).text = row[c] if c < len(row) else ""

        pr.report(90, "写出 DOCX")
        out.save(output)
    finally:
        doc.close()


def _word_text(
    src_path: str,
    output: str,
    pages: Optional[Sequence[int]],
    password: Optional[str],
    pr: Progress,
) -> None:
    import docx
    from docx.shared import Pt

    doc = _open(src_path, password)
    try:
        targets = list(pages) if pages else list(range(doc.page_count))
        out = docx.Document()
        style = out.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        total = len(targets)
        for k, i in enumerate(targets):
            pr.step(k, total, f"第 {i + 1} 页")
            if not 0 <= i < doc.page_count:
                continue
            if k:
                out.add_page_break()
            for para in _paragraphs(doc.load_page(i)):
                out.add_paragraph(para)
        pr.report(90, "写出 DOCX")
        out.save(output)
    finally:
        doc.close()


def _paragraphs(page: pymupdf.Page) -> List[str]:
    """把页面文字按视觉行聚合成段落，CJK 之间不插空格。"""
    data = page.get_text("dict")
    lines: List[Dict[str, Any]] = []
    for block in data.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            text = "".join(s.get("text", "") for s in line.get("spans", []))
            if text.strip():
                bbox = line.get("bbox") or (0, 0, 0, 0)
                lines.append({"text": text, "bbox": bbox})

    lines.sort(key=lambda l: (round(l["bbox"][1], 1), l["bbox"][0]))

    paragraphs: List[str] = []
    current = ""
    last_bottom = None
    for line in lines:
        text = line["text"].strip()
        gap = line["bbox"][1] - last_bottom if last_bottom is not None else 999
        if current and gap > 8:
            paragraphs.append(current)
            current = ""
        current = _join_cjk(current, text)
        last_bottom = line["bbox"][3]
    if current:
        paragraphs.append(current)
    return paragraphs


_CJK = re.compile(r"[㐀-䶿一-豈-]")


def _join_cjk(prev: str, nxt: str) -> str:
    if not prev:
        return nxt
    if _CJK.search(prev[-1]) and _CJK.search(nxt[0]):
        return prev + nxt
    return prev + " " + nxt


def _word_fidelity(
    src_path: str,
    output: str,
    pages: Optional[Sequence[int]],
    password: Optional[str],
    pr: Progress,
) -> None:
    import docx
    from docx.shared import Inches, Pt

    doc = _open(src_path, password)
    try:
        targets = list(pages) if pages else list(range(doc.page_count))
        out = docx.Document()
        for section in out.sections:
            section.top_margin = Pt(0)
            section.bottom_margin = Pt(0)
            section.left_margin = Pt(0)
            section.right_margin = Pt(0)

        total = len(targets)
        for k, i in enumerate(targets):
            pr.step(k, total, f"渲染第 {i + 1} 页")
            if not 0 <= i < doc.page_count:
                continue
            page = doc.load_page(i)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            buf = io.BytesIO(pix.tobytes("png"))

            if k:
                out.add_page_break()
            width_in = page.rect.width / 72.0
            out.add_picture(buf, width=Inches(width_in))
        pr.report(90, "写出 DOCX")
        out.save(output)
    finally:
        doc.close()


# ====================================================================== Excel


@dataclass
class TableData:
    page: int
    index: int
    rows: List[List[str]] = field(default_factory=list)
    method: str = "ruled"
    selected: bool = True

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def col_count(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def detect_tables(
    src_path: str,
    pages: Optional[Sequence[int]] = None,
    password: Optional[str] = None,
    progress: Progress_ = None,
) -> List[TableData]:
    """检测 PDF 中的表格（矢量线框法 + 对齐法）。"""
    pr = progress or Progress()
    doc = _open(src_path, password)
    try:
        targets = list(pages) if pages else list(range(doc.page_count))
        out: List[TableData] = []
        counter = 0
        for k, i in enumerate(targets):
            pr.step(k, len(targets), f"检测第 {i + 1} 页")
            if not 0 <= i < doc.page_count:
                continue
            page = doc.load_page(i)
            found = page.find_tables()
            for t in found.tables or []:
                rows = [[_clean_cell(c) for c in row] for row in (t.extract() or [])]
                rows = [r for r in rows if any(c.strip() for c in r)]
                if len(rows) < 2:
                    continue
                counter += 1
                out.append(TableData(page=i, index=counter, rows=rows, method="ruled"))

            if not (found.tables or []):
                aligned = _aligned_table(page)
                if aligned:
                    counter += 1
                    out.append(
                        TableData(page=i, index=counter, rows=aligned, method="aligned")
                    )
        pr.report(100, "完成")
        return out
    finally:
        doc.close()


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s*\n\s*", "\n", str(value)).strip()


def _aligned_table(page: pymupdf.Page) -> List[List[str]]:
    """无边框表格：按基线聚行、按 x 间隙聚列。"""
    words = page.get_text("words") or []
    if not words:
        return []

    rows: Dict[int, List] = {}
    for w in words:
        key = round(w[1] / 3)
        rows.setdefault(key, []).append(w)

    ordered = [rows[k] for k in sorted(rows)]
    if len(ordered) < 2:
        return []

    # 列边界：所有行 x0 聚类
    xs: List[float] = []
    for group in ordered:
        for w in group:
            xs.append(w[0])
    xs.sort()
    cols: List[float] = []
    for x in xs:
        if not cols or x - cols[-1] > 12:
            cols.append(x)
    if len(cols) < 2:
        return []

    table: List[List[str]] = []
    for group in ordered:
        group.sort(key=lambda w: w[0])
        cells = [""] * len(cols)
        for w in group:
            idx = max(range(len(cols)), key=lambda c: cols[c] if cols[c] <= w[0] + 4 else -1)
            cells[idx] = (cells[idx] + " " + w[4]).strip()
        if any(c for c in cells):
            table.append(cells)
    return table if len(table) >= 2 else []


def _parse_number(text: str) -> Optional[float]:
    t = (text or "").strip().replace(",", "").replace("，", "")
    t = t.replace("$", "").replace("¥", "").replace("%", "")
    if not t or not re.match(r"^[+-]?\d+(\.\d+)?$", t):
        return None
    try:
        return float(t)
    except ValueError:
        return None


def export_tables(
    tables: Sequence[TableData],
    output: str,
    fmt: str = EXCEL_XLSX,
    sheet_mode: str = SHEET_PER_TABLE,
    numbers: bool = True,
    auto_width: bool = True,
    progress: Progress_ = None,
) -> str:
    pr = progress or Progress()
    selected = [t for t in tables if t.selected and t.rows]
    if not selected:
        raise ValueError("没有选中任何表格")
    _safe_dir(output)

    if fmt == EXCEL_CSV:
        _write_csv(selected, output, pr)
    else:
        _write_xlsx(selected, output, sheet_mode, numbers, auto_width, pr)

    pr.report(100, "完成")
    return output


def _write_csv(tables: Sequence[TableData], output: str, pr: Progress) -> None:
    blocks: List[str] = []
    for k, t in enumerate(tables):
        pr.step(k, len(tables), f"表格 {t.index}")
        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\r\n")
        for row in t.rows:
            writer.writerow(row)
        blocks.append(buf.getvalue().rstrip("\r\n"))
    with open(output, "w", encoding="utf-8-sig", newline="") as fh:
        fh.write("\r\n\r\n".join(blocks))


def _write_xlsx(
    tables: Sequence[TableData],
    output: str,
    sheet_mode: str,
    numbers: bool,
    auto_width: bool,
    pr: Progress,
) -> None:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)

    def _new_sheet(name: str):
        safe = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
        base, i = safe, 1
        while safe in wb.sheetnames:
            i += 1
            safe = f"{base[:28]}_{i}"
        return wb.create_sheet(safe)

    def _write_rows(ws, rows, start_row: int) -> int:
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = ws.cell(row=start_row + r, column=c + 1)
                if numbers:
                    num = _parse_number(value)
                    cell.value = num if num is not None else value
                else:
                    cell.value = value
        return start_row + len(rows)

    def _autofit(ws) -> None:
        if not auto_width:
            return
        widths: Dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                length = len(str(cell.value))
                widths[cell.column] = max(widths.get(cell.column, 8), min(60, length + 2))
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = w

    if sheet_mode == SHEET_PER_TABLE:
        for k, t in enumerate(tables):
            pr.step(k, len(tables), f"表格 {t.index}")
            ws = _new_sheet(f"Table {t.index}")
            _write_rows(ws, t.rows, 1)
            _autofit(ws)
    elif sheet_mode == SHEET_PER_PAGE:
        pages = sorted({t.page for t in tables})
        for k, pno in enumerate(pages):
            pr.step(k, len(pages), f"第 {pno + 1} 页")
            ws = _new_sheet(f"Page {pno + 1}")
            row = 1
            for t in [t for t in tables if t.page == pno]:
                row = _write_rows(ws, t.rows, row) + 1
            _autofit(ws)
    else:
        ws = _new_sheet("Sheet1")
        row = 1
        for k, t in enumerate(tables):
            pr.step(k, len(tables), f"表格 {t.index}")
            row = _write_rows(ws, t.rows, row) + 1
        _autofit(ws)

    wb.save(output)
